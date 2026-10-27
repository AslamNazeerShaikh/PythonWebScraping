"""Role-targeted LinkedIn email collector (spec: docs/SPEC-linkedin-role-email-collector.md).

This is a **separate** path from the broad ``collect-emails`` collector. It
searches LinkedIn content for ``"<phrase> @<company-domain>"``, reads the
Latest feed, expands post bodies, and keeps only emails on the exact company
domain in a post section that explicitly names the phrase.

It deliberately does not write SQLite: the workbook is the artifact. The CLI
checkpoints after every domain and can resume into the same workbook. The
existing ``collect-emails`` command, its schema, and its dedupe rules are
untouched.
"""
from __future__ import annotations

import logging
import re
import unicodedata
from collections import OrderedDict
from collections.abc import Callable
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from playwright.sync_api import BrowserContext, Page
from pydantic import BaseModel, field_validator

from .config import SETTINGS
from .linkedin_email_scraper import (_check_bot_wall, _count_posts,
                                     _expand_show_more, _on_login_wall,
                                     _open_feed, _scroll_to_bottom,
                                     _wheel_feed)
from .models import EMAIL_RE
from .stealth_browser import human_pause

log = logging.getLogger(__name__)

CONTENT_SEARCH = "https://www.linkedin.com/search/results/content/"
DEFAULT_PHRASE = "asp.net core"
DEFAULT_DOMAINS = (
    "ltm.com", "capgemini.com", "tcs.com", "ey.com", "accenture.com",
    "deloitte.com", "harman.com", "altimetrik.com", "wipro.com",
    "samsung.com", "kaseya.com", "pwc.com", "hexaware.com",
    "persistent.com", "ascendion.com",
)
COMPANY_NAMES = {
    "ltm.com": "LTIMindtree",
    "capgemini.com": "Capgemini",
    "tcs.com": "TCS",
    "ey.com": "EY",
    "accenture.com": "Accenture",
    "deloitte.com": "Deloitte",
    "harman.com": "HARMAN",
    "altimetrik.com": "Altimetrik",
    "wipro.com": "Wipro",
    "samsung.com": "Samsung",
    "kaseya.com": "Kaseya",
    "pwc.com": "PwC",
    "hexaware.com": "Hexaware",
    "persistent.com": "Persistent",
    "ascendion.com": "Ascendion",
}

TARGET_EMAILS = 10
CHECKPOINT_POSTS = 20
MAX_POSTS = 30
SEARCH_WAIT_MS = 4500
NAV_ATTEMPTS = 3
NAV_RETRY_WAIT_MS = 4000
EXTRACT_WAIT_MS = 1100
SCROLL_WAIT_MS = 2800
STALL_WAIT_MS = 4500
MAX_STALLS = 3
MAX_ROUNDS = 12
DEFAULT_PAUSE_S = 15
FAILED_STATUSES = frozenset({"nav_timeout", "wall", "login_wall"})

HASHTAG_RE = re.compile(r"#\w+")
SECTION_BREAK_RE = re.compile(r"\n\s*(?:[━═]{3,}|[-—]{5,})\s*\n")
HIRE_RE = re.compile(
    r"\b(hiring|looking for|looking to hire|requirement|requirements|"
    r"opening|openings|vacanc(?:y|ies)|position|role|share (?:your )?"
    r"(?:resume|cv|profile)|send (?:your )?(?:resume|cv|profile)|apply|"
    r"walk[-\s]?in|immediate joiner|join (?:our|the) team)\b", re.I)
ROLE_RE = re.compile(
    r"\b(developer|engineer|full[-\s]?stack|backend|software|programmer|"
    r"architect|lead|consultant)\b", re.I)
JUNK_TLDS = frozenset({"png", "jpg", "jpeg", "gif", "svg", "webp", "bmp", "ico"})

ROLE_EXTRACT_JS = r"""
() => {
  /*role-content-cards*/
  const cards = [];
  const seen = new Set();
  const lc = document.querySelector('[data-testid="lazy-column"]');
  if (!lc) return cards;
  const items = Array.from(lc.querySelectorAll('[role="listitem"]'))
    .filter((it) => it.querySelector('[data-testid="expandable-text-box"]'));
  for (const it of items) {
    const box = it.querySelector('[data-testid="expandable-text-box"]');
    const text = (box.innerText || '').trim();
    if (!text) continue;
    let author = '';
    for (const a of it.querySelectorAll('a[href*="/in/"]')) {
      const t = (a.innerText || '').trim();
      if (t && t.indexOf('\u2022') !== -1) {
        author = t.split('\u2022')[0].trim();
        break;
      }
    }
    if (!author) {
      const a = it.querySelector('a[href*="/in/"]');
      if (a) author = (a.innerText || '').trim().split('\n')[0].trim();
    }
    const lines = (it.innerText || '').split('\n').map(s => s.trim()).filter(Boolean);
    let headline = '';
    const bi = lines.findIndex(x => x.includes('\u2022') && /(1st|2nd|3rd)/.test(x));
    if (bi >= 0 && lines[bi + 1] &&
        !/^(\d+[smhdw]|\d+mo|\d+yr|just now|(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\s+\d{1,2})$/i.test(lines[bi + 1])) {
      headline = lines[bi + 1];
    }
    let posted = '';
    for (const line of lines) {
      if (/^(\d+[smhdw]|\d+mo|\d+yr|just now|(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\s+\d{1,2})$/i.test(line)) {
        posted = line;
        break;
      }
    }
    let postId = it.getAttribute('componentkey') ||
                 it.getAttribute('data-id') ||
                 it.getAttribute('data-urn') || '';
    const m = postId.match(/update-card-focus(.+?)FeedType_/);
    if (m) postId = m[1];
    if (!postId) postId = (author || '') + '|' + (posted || '') + '|' + text.slice(0, 80);
    if (seen.has(postId)) continue;
    seen.add(postId);
    let url = '';
    const link = it.querySelector('a[href*="/feed/update/"], a[href*="/posts/"], a[href*="urn:li:activity"]');
    if (link) url = link.href || '';
    cards.push({ text, post_id: postId, author, headline, posted, url });
  }
  return cards;
}
"""

EMAIL_COLUMNS = [
    "company", "company_domain", "email", "role", "location", "post_rank",
    "author", "author_headline", "posted", "post_reference", "post_id",
    "post_url", "phrase_evidence", "email_context", "post_excerpt",
    "search_query", "search_url", "collected_at",
]
EMAIL_LIST_COLUMNS = [
    "email", "company", "company_domain", "role", "location", "posted",
    "post_rank", "author", "post_reference",
]
POST_COLUMNS = [
    "company", "company_domain", "post_rank", "author", "author_headline",
    "posted", "post_id", "post_url", "phrase_match", "hiring", "role",
    "domain_emails", "qualified_emails", "post_excerpt",
]
SUMMARY_COLUMNS = [
    "company", "company_domain", "posts_scanned", "phrase_posts",
    "qualified_posts", "unique_emails", "target_met", "status", "search_url",
]


class RoleEmailContact(BaseModel):
    """One validated email and its post evidence."""

    email: str
    company: str
    company_domain: str
    post_id: str = ""
    post_url: str = ""
    post_reference: str = ""
    author: str = ""
    author_headline: str = ""
    posted: str = ""
    role: str = ""
    location: str = ""
    post_rank: int = 0
    phrase_evidence: str = ""
    email_context: str = ""
    post_excerpt: str = ""
    search_query: str = ""
    search_url: str = ""
    collected_at: str = ""

    @field_validator("email")
    @classmethod
    def _valid_email(cls, v: str) -> str:
        v = (v or "").strip().lower()
        if not EMAIL_RE.fullmatch(v):
            raise ValueError(f"invalid email address: {v!r}")
        return v

    @field_validator("company_domain")
    @classmethod
    def _valid_domain(cls, v: str) -> str:
        v = (v or "").strip().lower().lstrip("@")
        if "." not in v:
            raise ValueError(f"invalid company domain: {v!r}")
        return v

    @field_validator("author", "author_headline", "posted", "role", "location")
    @classmethod
    def _squeeze(cls, v: str) -> str:
        return " ".join((v or "").split())

    @field_validator("post_url")
    @classmethod
    def _strip_url(cls, v: str) -> str:
        return (v or "").strip()


def build_search_url(phrase: str, domain: str) -> str:
    """Exact LinkedIn content-search URL for ``<phrase> @<domain>``.

    The ``sortBy`` value is the native JSON array string
    ``["date_posted"]``, not the legacy quoted scalar.
    """
    keyword = f"{phrase.strip()} @{domain.strip().lower().lstrip('@')}"
    return ("https://www.linkedin.com/search/results/content/?"
            f"keywords={quote(keyword, safe='')}&origin=FACETED_SEARCH"
            "&sortBy=%5B%22date_posted%22%5D")


def parse_domains(value: str) -> list[str]:
    """Parse comma-separated domains, defaulting to :data:`DEFAULT_DOMAINS`."""
    raw = [d.strip().lower().lstrip("@")
           for d in (value or "").split(",") if d.strip()]
    out: list[str] = []
    for domain in raw:
        if domain not in out:
            out.append(domain)
    return out or list(DEFAULT_DOMAINS)


def clean_text(text: str) -> str:
    """Normalise Unicode/zero-width characters and deobfuscate email syntax."""
    text = unicodedata.normalize("NFKC", text or "")
    text = re.sub(r"[\u200b-\u200f\u202a-\u202e\u2060\ufeff]", "", text)
    text = text.replace("\u00a0", " ")
    text = re.sub(r"\s*\[\s*at\s*\]\s*", "@", text, flags=re.I)
    text = re.sub(r"\s*\(\s*at\s*\)\s*", "@", text, flags=re.I)
    text = re.sub(r"\s*\[\s*dot\s*\]\s*", ".", text, flags=re.I)
    text = re.sub(r"\s*\(\s*dot\s*\)\s*", ".", text, flags=re.I)
    return text


def masked_hashtags(text: str) -> str:
    """Blank hashtags so ``#ASPNETCore`` cannot satisfy the phrase gate."""
    return HASHTAG_RE.sub(lambda m: " " * len(m.group(0)), text)


def phrase_pattern(phrase: str) -> re.Pattern[str]:
    """Build a whitespace/punctuation-tolerant regex for the role phrase.

    Dot-separated tokens are flexible: ``asp.net core`` also matches
    ``ASPNET Core``, ``ASP.NETCore`` and ``asp net core``.
    """
    def _token(token: str) -> str:
        parts = token.split(".")
        return r"\s*\.?\s*".join(re.escape(part) for part in parts)

    tokens = [_token(t) for t in phrase.split() if t]
    body = r"[\s._-]*".join(tokens) if tokens else re.escape(phrase)
    return re.compile(rf"(?<![a-z#]){body}(?![a-z])", re.I)


def phrase_matches(text: str, phrase: str) -> list[re.Match[str]]:
    """Return non-hashtag phrase matches in cleaned text."""
    return list(phrase_pattern(phrase).finditer(masked_hashtags(clean_text(text))))


def section_breaks(text: str) -> list[int]:
    """Positions of strong multi-role separators in ``text``."""
    return [m.start() for m in SECTION_BREAK_RE.finditer(text)]


def same_section(text: str, first: int, second: int) -> bool:
    """True when no strong separator sits between two text positions."""
    lo, hi = sorted((first, second))
    return not any(lo < pos < hi for pos in section_breaks(text))


def domain_email_matches(text: str, domain: str) -> list[tuple[str, int, int]]:
    """Exact-domain emails with positions, tolerating spaces around @ and dots."""
    target = domain.strip().lower().lstrip("@")
    domain_re = r"\s*\.\s*".join(re.escape(part) for part in target.split("."))
    pattern = re.compile(
        rf"([a-z0-9._%+\-]+)\s*@\s*{domain_re}(?![a-z0-9.\-])", re.I)
    out: list[tuple[str, int, int]] = []
    seen: set[str] = set()
    for match in pattern.finditer(clean_text(text)):
        local = match.group(1).strip(".-_")
        if not local:
            continue
        email = f"{local}@{target}".lower()
        if email.rsplit(".", 1)[-1] in JUNK_TLDS or email in seen:
            continue
        seen.add(email)
        out.append((email, match.start(), match.end()))
    return out


def phrase_context(text: str, phrase: str, radius: int = 200) -> str:
    """Context around the first real phrase match."""
    match = next(iter(phrase_matches(text, phrase)), None)
    if match is None:
        return ""
    return text[max(0, match.start() - radius): match.end() + radius]


def email_context(text: str, email: str, radius: int = 180) -> str:
    """Context around ``email`` in ``text``."""
    idx = text.lower().find(email.lower())
    if idx < 0:
        return ""
    return text[max(0, idx - radius): idx + len(email) + radius]


def role_hint(text: str) -> str:
    """Best-effort role/title line for the workbook."""
    lines = [" ".join(line.split()) for line in text.splitlines() if line.strip()]
    for line in lines:
        match = re.match(
            r"(?i)^(?:role|position|job title|we(?:'re| are) hiring|hiring)"
            r"\s*[:\-–]?\s*(.+)$", line)
        if match and not re.match(r"(?i)^summary\b", match.group(1)):
            return match.group(1)[:200]
    for line in lines:
        if (re.search(r"(?i)\b(developer|engineer|architect|lead|consultant)\b", line)
                and not re.search(r"(?i)\b(skills?|experience|location)\b", line)):
            return line[:200]
    return lines[0][:200] if lines else ""


def location_hint(text: str) -> str:
    """Best-effort location line for the workbook."""
    for line in text.splitlines():
        if re.search(r"(?i)\blocation\b|📍", line):
            cleaned = re.sub(r"(?i)^.*?(?:\blocation\b|📍)\s*[:\-]?\s*", "", line)
            return " ".join(cleaned.split())[:120]
    return ""


def clean_post_id(value: object) -> str:
    """Strip LinkedIn's wrapper around an opaque content-search post id."""
    text = str(value or "").strip()
    if text.lower() in {"", "nan", "none"}:
        return ""
    text = re.sub(r"^expanded", "", text)
    return re.sub(r"FeedType_.*$", "", text)


def build_contacts(cards: list[dict], company: str, domain: str, phrase: str,
                   collected_at: str | None = None) -> list[RoleEmailContact]:
    """Turn raw cards into validated, section-checked contacts."""
    stamp = collected_at or datetime.now().isoformat(timespec="seconds")
    query = f"{phrase.strip()} @{domain}"
    url = build_search_url(phrase, domain)
    out: list[RoleEmailContact] = []
    for rank, card in enumerate(cards, start=1):
        text = clean_text(str(card.get("text") or ""))
        if not text or not phrase_matches(text, phrase):
            continue
        if not HIRE_RE.search(text) or not ROLE_RE.search(text):
            continue
        emails = domain_email_matches(text, domain)
        if not emails:
            continue
        phrase_spans = [m.start() for m in phrase_matches(text, phrase)]
        post_id = clean_post_id(card.get("post_id"))
        post_url = str(card.get("url") or "").strip()
        post_ref = post_url or post_id
        for email, start, _end in emails:
            if not any(same_section(text, span, start) for span in phrase_spans):
                continue
            out.append(RoleEmailContact(
                email=email,
                company=company,
                company_domain=domain,
                post_id=post_id,
                post_url=post_url,
                post_reference=post_ref,
                author=str(card.get("author") or ""),
                author_headline=str(card.get("headline") or ""),
                posted=str(card.get("posted") or ""),
                role=role_hint(text),
                location=location_hint(text),
                post_rank=rank,
                phrase_evidence=phrase_context(text, phrase),
                email_context=email_context(text, email),
                post_excerpt=text[:3000],
                search_query=query,
                search_url=url,
                collected_at=stamp,
            ))
    return out


def extract_cards(page: Page) -> list[dict]:
    """Evaluate the SDUI card extractor; never raise on a detached page."""
    try:
        return page.evaluate(ROLE_EXTRACT_JS) or []
    except Exception:
        return []


def merge_cards(target: OrderedDict[str, dict], cards: list[dict]) -> None:
    """Append unseen cards in feed order."""
    for card in cards:
        post_id = str(card.get("post_id") or "")
        if not post_id or post_id in target:
            continue
        card = dict(card)
        card["post_id"] = post_id
        card["text"] = clean_text(str(card.get("text") or ""))
        target[post_id] = card


def open_search(page: Page, url: str) -> bool:
    """Navigate to a content search with a commit-level retry fallback."""
    for attempt in range(1, NAV_ATTEMPTS + 1):
        try:
            page.goto(url, wait_until="commit")
            page.wait_for_timeout(SEARCH_WAIT_MS)
            try:
                page.wait_for_selector('[data-testid="lazy-column"]', timeout=12000)
            except Exception:
                pass
            return True
        except Exception as exc:
            log.warning("Search navigation failed (%d/%d): %s",
                        attempt, NAV_ATTEMPTS, exc)
            try:
                page.evaluate("window.stop()")
            except Exception:
                pass
            page.wait_for_timeout(NAV_RETRY_WAIT_MS)
    return False


def build_post_rows(posts: list[dict], contacts: list[RoleEmailContact],
                    company: str, domain: str, phrase: str) -> list[dict]:
    """Build the audit rows for every scanned post."""
    qualified_by_post: dict[str, list[str]] = {}
    for contact in contacts:
        qualified_by_post.setdefault(contact.post_id, []).append(contact.email)
    rows: list[dict] = []
    for rank, post in enumerate(posts, start=1):
        text = clean_text(str(post.get("text") or ""))
        rows.append({
            "company": company,
            "company_domain": domain,
            "post_rank": rank,
            "author": str(post.get("author") or ""),
            "author_headline": str(post.get("headline") or ""),
            "posted": str(post.get("posted") or ""),
            "post_id": clean_post_id(post.get("post_id")),
            "post_url": str(post.get("url") or ""),
            "phrase_match": bool(phrase_matches(text, phrase)),
            "hiring": bool(HIRE_RE.search(text)),
            "role": bool(ROLE_RE.search(text)),
            "domain_emails": "; ".join(
                email for email, _start, _end in domain_email_matches(text, domain)),
            "qualified_emails": "; ".join(qualified_by_post.get(
                str(post.get("post_id") or ""), [])),
            "post_excerpt": text[:3000],
        })
    return rows


def collect_domain(page: Page, company: str, domain: str, phrase: str,
                   target: int = TARGET_EMAILS,
                   checkpoint_posts: int = CHECKPOINT_POSTS,
                   max_posts: int = MAX_POSTS
                   ) -> tuple[list[RoleEmailContact], list[dict], str]:
    """Search one company domain and return contacts, posts, and status."""
    url = build_search_url(phrase, domain)
    if not open_search(page, url):
        return [], [], "nav_timeout"
    if _on_login_wall(page):
        return [], [], "login_wall"
    if _check_bot_wall(page):
        return [], [], "wall"

    posts: OrderedDict[str, dict] = OrderedDict()
    stalls = 0
    for round_no in range(MAX_ROUNDS):
        _expand_show_more(page)
        page.wait_for_timeout(EXTRACT_WAIT_MS)
        merge_cards(posts, extract_cards(page))
        current = list(posts.values())[:max_posts]
        contacts = build_contacts(current, company, domain, phrase)
        unique = {contact.email for contact in contacts}
        if len(current) >= max_posts:
            break
        if len(current) >= checkpoint_posts and len(unique) >= target:
            break
        if round_no >= MAX_ROUNDS - 1:
            break
        before = _count_posts(page)
        _wheel_feed(page)
        page.wait_for_timeout(SCROLL_WAIT_MS)
        after = _count_posts(page)
        if after <= before:
            _scroll_to_bottom(page)
            page.wait_for_timeout(STALL_WAIT_MS)
            after = max(after, _count_posts(page))
            stalls = stalls + 1 if after <= before else 0
        else:
            stalls = 0
        if stalls >= MAX_STALLS:
            break

    current = list(posts.values())[:max_posts]
    contacts = build_contacts(current, company, domain, phrase)
    deduped: list[RoleEmailContact] = []
    seen: set[str] = set()
    for contact in contacts:
        if contact.email in seen:
            continue
        seen.add(contact.email)
        deduped.append(contact)
    if not current:
        status = "no_results"
    elif len(deduped) < target:
        status = "below_target"
    else:
        status = "ok"
    return deduped, current, status


def run_role_collect(
        context: BrowserContext,
        domains: list[str],
        phrase: str = DEFAULT_PHRASE,
        target: int = TARGET_EMAILS,
        checkpoint_posts: int = CHECKPOINT_POSTS,
        max_posts: int = MAX_POSTS,
        pause_s: int = DEFAULT_PAUSE_S,
        on_domain: Callable[[list[RoleEmailContact], list[dict], list[dict]], None]
        | None = None,
        initial_contacts: list[RoleEmailContact] | None = None,
        initial_posts: list[dict] | None = None,
        initial_summary: list[dict] | None = None,
) -> tuple[list[RoleEmailContact], list[dict], list[dict]]:
    """Run the role collector across ``domains`` using one browser page.

    ``initial_*`` rows make a run resumable: completed domains stay in the
    workbook, a rerun replaces that domain's rows, and a failed rerun keeps the
    previous rows rather than erasing good evidence.
    """
    page = context.new_page()
    if not _open_feed(page):
        return list(initial_contacts or []), list(initial_posts or []), \
            list(initial_summary or [])
    all_contacts: list[RoleEmailContact] = list(initial_contacts or [])
    all_posts: list[dict] = list(initial_posts or [])
    summary: list[dict] = list(initial_summary or [])
    for index, domain in enumerate(domains):
        company = COMPANY_NAMES.get(domain, domain)
        contacts, posts, status = collect_domain(
            page, company, domain, phrase, target, checkpoint_posts, max_posts)
        post_rows = build_post_rows(posts, contacts, company, domain, phrase)
        has_previous = (
            any(c.company_domain == domain for c in all_contacts)
            or any(row.get("company_domain") == domain for row in summary))
        if status in FAILED_STATUSES and has_previous:
            log.warning("%s failed (%s); keeping previous checkpoint rows.",
                        domain, status)
        else:
            all_contacts = [c for c in all_contacts
                            if c.company_domain != domain]
            all_posts = [row for row in all_posts
                         if row.get("company_domain") != domain]
            summary = [row for row in summary
                       if row.get("company_domain") != domain]
            all_contacts.extend(contacts)
            all_posts.extend(post_rows)
            summary.append({
                "company": company,
                "company_domain": domain,
                "posts_scanned": len(posts),
                "phrase_posts": sum(1 for row in post_rows
                                    if row["phrase_match"]),
                "qualified_posts": sum(
                    1 for row in post_rows if row["qualified_emails"]),
                "unique_emails": len(contacts),
                "target_met": len(contacts) >= target,
                "status": status,
                "search_url": build_search_url(phrase, domain),
            })
        if on_domain is not None:
            on_domain(all_contacts, all_posts, summary)
        if status in {"wall", "login_wall"}:
            log.warning("LinkedIn wall/login at %s; stopping with partial data.",
                        domain)
            break
        if status == "nav_timeout":
            log.warning("Navigation timed out for %s; continuing.", domain)
        if index < len(domains) - 1:
            human_pause(pause_s, pause_s + 5)
    return all_contacts, all_posts, summary


def _method_rows(phrase: str, target: int, checkpoint_posts: int,
                 max_posts: int) -> list[dict]:
    return [
        {"Item": "Collected", "Detail": datetime.now().strftime("%Y-%m-%d %H:%M IST")},
        {"Item": "Source", "Detail": "LinkedIn content search, signed-in browser session"},
        {"Item": "Query per domain", "Detail": f"{phrase} @<company-domain>"},
        {"Item": "Sort", "Detail": 'Latest (sortBy=["date_posted"])'},
        {"Item": "Date filter", "Detail": "None - latest matching posts"},
        {"Item": "Posts per domain",
         "Detail": f"Stop at {checkpoint_posts} if {target} emails reached; else max {max_posts}"},
        {"Item": "Post gate", "Detail": "Explicit non-hashtag phrase mention in the body"},
        {"Item": "Email gate", "Detail": "Only emails on the exact searched company domain"},
        {"Item": "Section gate", "Detail": "Email and phrase must be in the same multi-role section"},
        {"Item": "Location", "Detail": "No location filter applied"},
        {"Item": "Post reference", "Detail": "Post URL when exposed, otherwise opaque post id"},
        {"Item": "No evasion", "Detail": "No CAPTCHA solving, proxy rotation, or fingerprint spoofing"},
    ]


def read_role_workbook(path: Path | str
                       ) -> tuple[list[RoleEmailContact], list[dict], list[dict]]:
    """Read a prior role workbook for ``--resume``; empty when absent/corrupt."""
    source = Path(path)
    if not source.exists():
        return [], [], []
    try:
        email_df = pd.read_excel(source, sheet_name="Emails")
        post_df = pd.read_excel(source, sheet_name="Posts Scanned")
        summary_df = pd.read_excel(source, sheet_name="Summary")
    except Exception:
        return [], [], []
    contacts: list[RoleEmailContact] = []
    email_df = email_df.astype(object).where(pd.notna(email_df), "")
    for row in email_df.to_dict("records"):
        data = {key: row.get(key, "")
                for key in RoleEmailContact.model_fields}
        try:
            data["post_rank"] = int(float(data.get("post_rank") or 0))
            contacts.append(RoleEmailContact(**data))
        except Exception:
            continue
    clean = lambda df: df.astype(object).where(pd.notna(df), "").to_dict("records")
    return contacts, clean(post_df), clean(summary_df)


def write_role_workbook(path: Path | str, contacts: list[RoleEmailContact],
                        posts: list[dict], summary: list[dict],
                        phrase: str = DEFAULT_PHRASE,
                        target: int = TARGET_EMAILS,
                        checkpoint_posts: int = CHECKPOINT_POSTS,
                        max_posts: int = MAX_POSTS) -> Path:
    """Write the five-sheet role-email workbook and return its path."""
    dest = Path(path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    email_rows = [contact.model_dump() for contact in contacts]
    email_df = pd.DataFrame(email_rows, columns=EMAIL_COLUMNS)
    list_df = email_df[EMAIL_LIST_COLUMNS].copy() if not email_df.empty else \
        pd.DataFrame(columns=EMAIL_LIST_COLUMNS)
    post_df = pd.DataFrame(posts, columns=POST_COLUMNS)
    summary_df = pd.DataFrame(summary, columns=SUMMARY_COLUMNS)
    method_df = pd.DataFrame(_method_rows(phrase, target, checkpoint_posts,
                                          max_posts), columns=["Item", "Detail"])
    with pd.ExcelWriter(dest, engine="openpyxl") as writer:
        list_df.to_excel(writer, index=False, sheet_name="Email List")
        email_df.to_excel(writer, index=False, sheet_name="Emails")
        summary_df.to_excel(writer, index=False, sheet_name="Summary")
        post_df.to_excel(writer, index=False, sheet_name="Posts Scanned")
        method_df.to_excel(writer, index=False, sheet_name="Method")
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(vertical="center")
            for column in sheet.columns:
                letter = get_column_letter(column[0].column)
                width = min(58, max(12, max(
                    len(str(cell.value or "")) for cell in column) + 2))
                sheet.column_dimensions[letter].width = width
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
    return dest
